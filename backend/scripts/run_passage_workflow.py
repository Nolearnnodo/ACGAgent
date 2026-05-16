"""Batch entrypoint for the passage ingestion workflow.

Examples, from the repository root or backend/:

    # Dry run a directory of .txt files. Neo4j writes are stubbed.
    python -m scripts.run_passage_workflow test_data/ --dry-run

    # Dry run the micro-dataset manifest in the parent directory.
    python -m scripts.run_passage_workflow ../开发阶段微数据集.xlsx --corpus-root data/processed --dry-run --limit 2

    # Real run: writes Neo4j and uses the configured LLM provider.
    python -m scripts.run_passage_workflow ../开发阶段微数据集.xlsx --corpus-root data/processed

Output: one YAML-like text file per passage under output/.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from unittest.mock import MagicMock


def _patch_graph_repository_to_stub() -> None:
    """Replace GraphRepository with a noop stub for dry-run execution."""

    import app.graph.repository as repo_module

    def _make_stub():
        stub = MagicMock()
        success = {"status": "success", "operation": "stub"}
        for method in [
            "upsert_passage_info",
            "upsert_person",
            "upsert_life_event",
            "upsert_historical_event_time_link",
            "upsert_person_historical_event",
            "upsert_person_relation",
            "merge_same_name_persons_for_passage",
            "merge_person_nodes",
            "run_read_query",
            "run_write_query",
        ]:
            getattr(stub, method).return_value = success
        stub.merge_same_name_persons_for_passage.return_value = {
            "status": "success",
            "doc_id": 0,
            "matched_person_count": 0,
            "merged_person_count": 0,
            "failed_merge_count": 0,
            "merges": [],
            "failures": [],
        }
        return stub

    repo_module.GraphRepository = _make_stub
    for skill_module in (
        "app.skills.atomic.passage_meta",
        "app.skills.atomic.person_layer",
        "app.skills.atomic.event_relation",
        "app.skills.atomic.person_exact_match_merge",
    ):
        mod = __import__(skill_module, fromlist=["GraphRepository"])
        mod.GraphRepository = _make_stub  # type: ignore[attr-defined]


def _patch_sqlite_dictionary_to_stub() -> None:
    """Avoid requiring migrated SQLite dictionary tables during dry runs."""

    from app.services import dictionary_service as ds

    ds.reset_cache()
    ds._cache["era_entries"] = [
        ds.EraEntry(
            era="Kaiyuan",
            dynasty="Tang",
            start_year=713,
            end_year=741,
            simplified="Kaiyuan",
            traditional="Kaiyuan",
        )
    ]
    ds._cache["historical_events"] = []
    ds._cache["source_types"] = {0: "epitaph", 1: "history"}
    ds._cache["relation_codes"] = {
        code: ds.RelationCodeEntry(code=code, meaning=code, direction_hint="")
        for code in ("F", "M", "S", "D", "H", "W", "Z", "C", "B", "O")
    }

    class _NoopSession:
        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc, tb):
            return False

        def get(self, model, key):
            return None

        def add(self, item):
            return None

        def commit(self):
            return None

    passage_meta = __import__("app.skills.atomic.passage_meta", fromlist=["SessionLocal"])
    passage_meta.SessionLocal = lambda: _NoopSession()  # type: ignore[attr-defined]


def _read_xlsx_manifest_titles(manifest_path: Path, sheet_name: str | None = None) -> list[str]:
    """Read title strings from the first sheet, preserving manifest order."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read .xlsx manifests") from exc

    workbook = load_workbook(manifest_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        titles: list[str] = []
        seen: set[str] = set()
        for row in worksheet.iter_rows(values_only=True):
            for value in row:
                if not isinstance(value, str):
                    continue
                title = value.strip()
                if not title or title in seen:
                    continue
                seen.add(title)
                titles.append(title)
        return titles
    finally:
        workbook.close()


def _normalise_title_key(value: str) -> str:
    return "".join(str(value).replace("\u7b2c", "").split())


def _resolve_manifest_files(
    manifest_path: Path,
    corpus_root: Path,
    sheet_name: str | None = None,
) -> tuple[list[Path], list[str]]:
    titles = _read_xlsx_manifest_titles(manifest_path, sheet_name=sheet_name)

    stem_index: dict[str, Path] = {}
    normalised_index: dict[str, Path] = {}
    for path in sorted(corpus_root.rglob("*.txt")):
        stem_index.setdefault(path.stem, path)
        normalised_index.setdefault(_normalise_title_key(path.stem), path)

    files: list[Path] = []
    missing: list[str] = []
    for title in titles:
        path = stem_index.get(title) or normalised_index.get(_normalise_title_key(title))
        if path is None:
            missing.append(title)
        else:
            files.append(path)
    return files, missing


def _collect_input_files(args: argparse.Namespace, repo_root: Path) -> list[Path]:
    input_path = Path(args.input_path).resolve()
    if input_path.suffix.lower() == ".xlsx":
        corpus_root = (
            Path(args.corpus_root).resolve()
            if args.corpus_root
            else repo_root / "data" / "processed"
        )
        files, missing = _resolve_manifest_files(
            input_path,
            corpus_root=corpus_root,
            sheet_name=args.xlsx_sheet,
        )
        print(f"manifest: {input_path}")
        print(f"corpus root: {corpus_root}")
        print(f"matched {len(files)} titles, missing {len(missing)}")
        for title in missing[:10]:
            print(f"  missing: {title}")
        if len(missing) > 10:
            print(f"  ... {len(missing) - 10} more missing titles")
    elif input_path.is_file():
        files = [input_path]
    else:
        files = sorted(input_path.rglob("*.txt"))

    if args.limit is not None:
        files = files[: args.limit]
    return files


def _build_context(repo_root: Path, file_path: Path, doc_id: int, output_dir: Path):
    from app.agents.context import ExecutionContext

    text = file_path.read_text(encoding="utf-8")
    try:
        rel = file_path.relative_to(repo_root).as_posix()
    except ValueError:
        rel = file_path.as_posix()
    return ExecutionContext(
        user={"id": 1, "role": "admin", "email": "admin@example.com"},
        conversation={},
        metadata={
            "trigger_type": "passage_upload",
            "passage_id": doc_id,
            "passage": {
                "doc_id": doc_id,
                "title": file_path.stem,
                "context": text,
                "source_type": "upload",
                "source_file": rel,
            },
            "output_dir": str(output_dir),
        },
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="Run the passage ingestion workflow in batch.")
    parser.add_argument("input_path", help="A .txt file, a directory of .txt files, or an .xlsx manifest.")
    parser.add_argument(
        "--corpus-root",
        default=None,
        help="Root directory for resolving titles from an .xlsx manifest. Default: data/processed.",
    )
    parser.add_argument(
        "--xlsx-sheet",
        default=None,
        help="Optional sheet name for .xlsx manifests. Default: first worksheet.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Run only the first N resolved passages.",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Output directory. Default: repository output/.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Stub Neo4j writes. LLM still follows .env configuration.",
    )
    parser.add_argument(
        "--doc-id-base",
        type=int,
        default=90000,
        help="Starting doc_id base for this batch. Default: 90000.",
    )
    args = parser.parse_args()

    repo_root = Path(__file__).resolve().parents[2]
    backend_root = repo_root / "backend"
    if str(backend_root) not in sys.path:
        sys.path.insert(0, str(backend_root))

    if args.dry_run:
        _patch_graph_repository_to_stub()
        _patch_sqlite_dictionary_to_stub()
        print("(dry-run) GraphRepository is stubbed; Neo4j will not be written.")

    from app.skills.workflow.passage_ingestion_workflow import (
        PassageIngestionWorkflowSkill,
    )

    output_dir = Path(args.output_dir).resolve() if args.output_dir else repo_root / "output"
    output_dir.mkdir(parents=True, exist_ok=True)

    files = _collect_input_files(args, repo_root)
    if not files:
        print(f"no input .txt files found for {args.input_path}")
        return 1

    print(f"running {len(files)} passages -> {output_dir}")
    skill = PassageIngestionWorkflowSkill()

    failures = 0
    for idx, fp in enumerate(files, 1):
        doc_id = args.doc_id_base + idx
        try:
            context = _build_context(repo_root, fp, doc_id, output_dir)
            result = skill.run(context, arguments={})
            print(
                f"  [{idx:03d}] {fp.name:<55} status={result['status']:<8} "
                f"warnings={result['warning_count']:<3} -> {Path(result['output_path']).name}"
            )
        except Exception as exc:
            failures += 1
            print(f"  [{idx:03d}] {fp.name} ERROR: {exc}")

    print()
    print(f"done. failures: {failures} / total: {len(files)}")
    return 0 if failures == 0 else 2


if __name__ == "__main__":
    sys.exit(main())
